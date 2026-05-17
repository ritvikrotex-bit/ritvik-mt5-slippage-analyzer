"""
Excel parser for MT5 broker history reports.

Finds the Deals section dynamically (no hardcoded row numbers), extracts
all data rows, and filters them to the rows that are valid candidates
for slippage analysis (Direction = out, comment matches TP/SL pattern).

Returns a structured dict with:
  - valid_deals : list of row dicts
  - skipped     : {reason: count} audit map
  - total_rows  : int (every row processed below the header)
  - header_row  : int (1-indexed Excel row where header was found)
  - sheet_name  : str
  - unmatched_comments : list of comment strings that looked like TP/SL but didn't parse
"""
from datetime import datetime
from typing import Dict, List, Tuple, Any

from openpyxl import load_workbook

from config import (
    REQUIRED_HEADER_TOKENS,
    MIN_HEADER_MATCHES,
    MUST_HAVE_TOKENS,
    VALID_DIRECTION,
)
from comment_parser import parse_trigger


def find_deals_header(ws) -> Tuple[int, Dict[str, int]]:
    """
    Scan rows of the worksheet looking for the Deals header.

    A row qualifies if:
      - it contains at least MIN_HEADER_MATCHES of REQUIRED_HEADER_TOKENS
      - AND it contains all MUST_HAVE_TOKENS (this excludes the Positions header,
        which has 'Position' instead of 'Deal' and no 'Direction')

    Returns (header_row_1indexed, {column_name_lowercase: column_index_0based})
    Raises ValueError if no qualifying row is found.
    """
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row:
            continue
        tokens = {str(c).strip().lower() for c in row if c is not None and str(c).strip()}
        if not MUST_HAVE_TOKENS.issubset(tokens):
            continue
        if len(REQUIRED_HEADER_TOKENS & tokens) < MIN_HEADER_MATCHES:
            continue
        col_map = {
            str(v).strip().lower(): idx
            for idx, v in enumerate(row)
            if v is not None and str(v).strip()
        }
        return i, col_map
    raise ValueError("Could not locate Deals header row in worksheet")


def _is_valid_timestamp(value: Any) -> bool:
    """True iff the value is parseable as an MT5-style timestamp."""
    if value is None:
        return False
    if isinstance(value, datetime):
        return True
    s = str(value).strip()
    if not s:
        return False
    for fmt in (
        "%Y.%m.%d %H:%M:%S.%f",
        "%Y.%m.%d %H:%M:%S",
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
    ):
        try:
            datetime.strptime(s, fmt)
            return True
        except ValueError:
            continue
    return False


def _safe_float(value: Any):
    """Return float or None — never raises."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "")
    if not s:
        return None
    # Same European-decimal handling as comment parser
    if "," in s and "." in s:
        if s.rfind(".") > s.rfind(","):
            s = s.replace(",", "")
        else:
            s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def extract_deals(filepath: str) -> Dict[str, Any]:
    """
    Main entry. Reads the workbook, locates the Deals section, returns
    structured data + full audit trail.
    """
    wb = load_workbook(filepath, read_only=True, data_only=True)

    chosen_ws = None
    header_idx = None
    col_map = None

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        try:
            header_idx, col_map = find_deals_header(ws)
            chosen_ws = ws
            break
        except ValueError:
            continue

    if chosen_ws is None:
        raise ValueError("No Deals section found in any sheet of the workbook")

    # Required columns
    required = ["time", "type", "direction", "volume", "price",
                "comment", "deal", "order", "symbol"]
    missing = [c for c in required if c not in col_map]
    if missing:
        raise ValueError(
            f"Header found at row {header_idx} but missing required columns: {missing}"
        )

    def col(row, name):
        idx = col_map.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    valid_deals: List[Dict[str, Any]] = []
    skipped: Dict[str, int] = {}
    unmatched_comments: List[str] = []
    total_rows = 0

    for row in chosen_ws.iter_rows(min_row=header_idx + 1, values_only=True):
        time_val = col(row, "time")

        # Stop conditions: a non-timestamp Time cell means we've reached the
        # summary block (Balance:, Total Net Profit:, etc.) — done.
        if not _is_valid_timestamp(time_val):
            # If the entire row is empty, also stop.
            if all(c is None or str(c).strip() == "" for c in row):
                break
            # If just the time is non-parseable but other content exists, this
            # is the summary footer — stop.
            break

        total_rows += 1

        direction = str(col(row, "direction") or "").strip().lower()
        if direction != VALID_DIRECTION:
            key = f"direction={direction!r}" if direction else "direction_empty"
            skipped[key] = skipped.get(key, 0) + 1
            continue

        comment = col(row, "comment")
        comment_str = str(comment) if comment is not None else ""
        trigger = parse_trigger(comment_str)

        if trigger is None:
            # If the comment LOOKS like it should be TP/SL but didn't parse,
            # capture it so the user can see new broker formats.
            low = comment_str.lower()
            if "tp" in low or "sl" in low or "t/p" in low or "s/l" in low:
                unmatched_comments.append(comment_str)
            skipped["no_tp_sl_comment"] = skipped.get("no_tp_sl_comment", 0) + 1
            continue

        type_v = str(col(row, "type") or "").strip().lower()
        if type_v not in ("buy", "sell"):
            skipped[f"unknown_type={type_v!r}"] = skipped.get(
                f"unknown_type={type_v!r}", 0
            ) + 1
            continue

        volume = _safe_float(col(row, "volume"))
        executed = _safe_float(col(row, "price"))
        if volume is None or executed is None:
            skipped["non_numeric_price_or_volume"] = skipped.get(
                "non_numeric_price_or_volume", 0
            ) + 1
            continue
        if volume <= 0:
            skipped["non_positive_volume"] = skipped.get("non_positive_volume", 0) + 1
            continue

        valid_deals.append({
            "time": time_val,
            "deal": col(row, "deal"),
            "order": col(row, "order"),
            "symbol": col(row, "symbol"),
            "type": type_v,
            "direction": direction,
            "volume": volume,
            "executed_price": executed,
            "trigger_type": trigger[0],   # "TP" or "SL"
            "trigger_price": trigger[1],  # float
            "comment": comment_str,
            "commission": _safe_float(col(row, "commission")),
            "fee": _safe_float(col(row, "fee")) if "fee" in col_map else None,
            "swap": _safe_float(col(row, "swap")) if "swap" in col_map else None,
            "profit": _safe_float(col(row, "profit")) if "profit" in col_map else None,
            "balance": _safe_float(col(row, "balance")) if "balance" in col_map else None,
        })

    return {
        "valid_deals": valid_deals,
        "skipped": skipped,
        "total_rows": total_rows,
        "header_row": header_idx,
        "sheet_name": chosen_ws.title,
        "unmatched_comments": unmatched_comments,
    }
