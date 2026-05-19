"""
Excel parser for the Orders section and the Deals-direction-in subset
of MT5 broker history reports.

Self-contained: does NOT import from files/parser.py. Helper functions
(_safe_float, _is_valid_timestamp, header detection) are local copies
so this module can be deleted in one piece without affecting the TP/SL
engine. See LOGIC_LOCK.md §10.

Public entry points:
  - extract_orders(filepath)     -> filled pending orders + audit
  - extract_in_deals(filepath)   -> opening deals (Direction == 'in') + audit
"""
from collections import Counter
from datetime import datetime
from typing import Any, Dict, Tuple

from openpyxl import load_workbook

from config import (
    # Deals-section header detection (read-only constants)
    REQUIRED_HEADER_TOKENS,
    MIN_HEADER_MATCHES,
    MUST_HAVE_TOKENS,
    # Orders-section header detection
    ORDERS_REQUIRED_HEADER_TOKENS,
    ORDERS_MIN_HEADER_MATCHES,
    ORDERS_MUST_HAVE_TOKENS,
    # Filtering
    VALID_ORDER_STATE,
    PENDING_ORDER_TYPES,
)


# ---------------------------------------------------------------------------
# Local helpers (intentionally duplicated from files/parser.py for isolation)
# ---------------------------------------------------------------------------
def _is_valid_timestamp(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, datetime):
        return True
    s = str(value).strip()
    if not s:
        return False
    for fmt in (
        "%Y.%m.%d %H:%M:%S.%f",
        "%Y.%m.%d %H:%M:%S",
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
    ):
        try:
            datetime.strptime(s, fmt)
            return True
        except ValueError:
            continue
    return False


def _safe_float(value: Any):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        if s.rfind(".") > s.rfind(","):
            s = s.replace(",", "")
        else:
            s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _parse_volume_pair(value):
    """
    Orders Volume cells look like '0.2 / 0.2' = 'requested / filled'.
    Returns (requested_volume, filled_volume) as floats, or (None, None).
    Accepts a single numeric value too (treats it as both).
    """
    if value is None:
        return (None, None)
    if isinstance(value, (int, float)):
        v = float(value)
        return (v, v)
    s = str(value).strip()
    if not s:
        return (None, None)
    if "/" in s:
        left, _, right = s.partition("/")
        return (_safe_float(left), _safe_float(right))
    v = _safe_float(s)
    return (v, v)


# ---------------------------------------------------------------------------
# Header detection
# ---------------------------------------------------------------------------
def _find_orders_header(ws) -> Tuple[int, Dict[str, int]]:
    """
    Scan rows looking for the Orders header. Disambiguated from
    Positions/Deals by ORDERS_MUST_HAVE_TOKENS = {state, open time}.
    """
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row:
            continue
        tokens = {str(c).strip().lower() for c in row if c is not None and str(c).strip()}
        if not ORDERS_MUST_HAVE_TOKENS.issubset(tokens):
            continue
        if len(ORDERS_REQUIRED_HEADER_TOKENS & tokens) < ORDERS_MIN_HEADER_MATCHES:
            continue
        col_map = {
            str(v).strip().lower(): idx
            for idx, v in enumerate(row)
            if v is not None and str(v).strip()
        }
        return i, col_map
    raise ValueError("Could not locate Orders header row in worksheet")


def _find_deals_header(ws) -> Tuple[int, Dict[str, int]]:
    """
    Local copy of the Deals-header detection. Same logic as the TP/SL
    parser but lives here so this module is self-contained.
    """
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row:
            continue
        tokens = {str(c).strip().lower() for c in row if c is not None and str(c).strip()}
        if not MUST_HAVE_TOKENS.issubset(tokens):
            continue
        if len(REQUIRED_HEADER_TOKENS & tokens) < MIN_HEADER_MATCHES:
            continue
        col_map = {
            str(v).strip().lower(): idx
            for idx, v in enumerate(row)
            if v is not None and str(v).strip()
        }
        return i, col_map
    raise ValueError("Could not locate Deals header row in worksheet")


# ---------------------------------------------------------------------------
# Public entry points
# ---------------------------------------------------------------------------
def extract_orders(filepath: str) -> Dict[str, Any]:
    """
    Locate the Orders section and return filled pending orders.

    Returns:
      {
        'filled_pendings'   : list of dicts (filled buy/sell limit/stop/stop_limit)
        'skipped'           : {reason: count} audit map
        'total_rows'        : int
        'header_row'        : int (1-indexed)
        'sheet_name'        : str
        'state_counts'      : dict of all observed states (diagnostics)
        'type_counts'       : dict of all observed types (diagnostics)
      }
    """
    wb = load_workbook(filepath, read_only=True, data_only=True)

    chosen_ws = None
    header_idx = None
    col_map = None

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        try:
            header_idx, col_map = _find_orders_header(ws)
            chosen_ws = ws
            break
        except ValueError:
            continue

    if chosen_ws is None:
        raise ValueError("No Orders section found in any sheet of the workbook")

    required = ["open time", "order", "symbol", "type", "volume",
                "price", "time", "state"]
    missing = [c for c in required if c not in col_map]
    if missing:
        raise ValueError(
            f"Orders header found at row {header_idx} but missing required columns: {missing}"
        )

    def col(row, name):
        idx = col_map.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    filled_pendings = []
    skipped: Dict[str, int] = {}
    state_counts: Counter = Counter()
    type_counts: Counter = Counter()
    total_rows = 0

    for row in chosen_ws.iter_rows(min_row=header_idx + 1, values_only=True):
        open_time = col(row, "open time")

        if not _is_valid_timestamp(open_time):
            # Stop at end-of-section / blank / "Deals" header / summary row.
            first_cell = row[0] if row else None
            if first_cell is None or str(first_cell).strip() == "" or \
                    str(first_cell).strip().lower() == "deals":
                break
            break

        total_rows += 1

        order_type_raw = str(col(row, "type") or "").strip().lower()
        state_raw = str(col(row, "state") or "").strip().lower()
        type_counts[order_type_raw] += 1
        state_counts[state_raw] += 1

        if state_raw != VALID_ORDER_STATE:
            key = f"state={state_raw!r}" if state_raw else "state_empty"
            skipped[key] = skipped.get(key, 0) + 1
            continue

        if order_type_raw not in PENDING_ORDER_TYPES:
            key = f"non_pending_type={order_type_raw!r}"
            skipped[key] = skipped.get(key, 0) + 1
            continue

        requested = _safe_float(col(row, "price"))
        if requested is None or requested <= 0:
            skipped["bad_requested_price"] = skipped.get("bad_requested_price", 0) + 1
            continue

        req_vol, filled_vol = _parse_volume_pair(col(row, "volume"))
        if filled_vol is None or filled_vol <= 0:
            skipped["bad_volume"] = skipped.get("bad_volume", 0) + 1
            continue

        order_id = col(row, "order")
        if order_id is None:
            skipped["missing_order_id"] = skipped.get("missing_order_id", 0) + 1
            continue

        filled_pendings.append({
            "open_time": open_time,
            "order_id": order_id,
            "symbol": col(row, "symbol"),
            "type": order_type_raw,         # "buy limit", "sell stop", etc.
            "requested_volume": req_vol,
            "filled_volume": filled_vol,
            "requested_price": requested,
            "sl": _safe_float(col(row, "s / l")) if "s / l" in col_map else None,
            "tp": _safe_float(col(row, "t / p")) if "t / p" in col_map else None,
            "fill_time": col(row, "time"),
            "state": state_raw,
            "comment": str(col(row, "comment") or "") if "comment" in col_map else "",
        })

    return {
        "filled_pendings": filled_pendings,
        "skipped": skipped,
        "total_rows": total_rows,
        "header_row": header_idx,
        "sheet_name": chosen_ws.title,
        "state_counts": dict(state_counts),
        "type_counts": dict(type_counts),
    }


def extract_in_deals(filepath: str) -> Dict[str, Any]:
    """
    Return Deals where Direction == 'in' — the opening legs of positions,
    needed to match against filled pending orders by Order ID.
    """
    wb = load_workbook(filepath, read_only=True, data_only=True)

    chosen_ws = None
    header_idx = None
    col_map = None

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        try:
            header_idx, col_map = _find_deals_header(ws)
            chosen_ws = ws
            break
        except ValueError:
            continue

    if chosen_ws is None:
        raise ValueError("No Deals section found in any sheet of the workbook")

    required = ["time", "deal", "order", "symbol", "type", "direction",
                "volume", "price"]
    missing = [c for c in required if c not in col_map]
    if missing:
        raise ValueError(
            f"Deals header found at row {header_idx} but missing required columns: {missing}"
        )

    def col(row, name):
        idx = col_map.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    in_deals = []
    skipped: Dict[str, int] = {}
    total_rows = 0

    for row in chosen_ws.iter_rows(min_row=header_idx + 1, values_only=True):
        time_val = col(row, "time")
        if not _is_valid_timestamp(time_val):
            if all(c is None or str(c).strip() == "" for c in row):
                break
            break

        total_rows += 1

        direction = str(col(row, "direction") or "").strip().lower()
        if direction != "in":
            key = f"direction={direction!r}" if direction else "direction_empty"
            skipped[key] = skipped.get(key, 0) + 1
            continue

        type_v = str(col(row, "type") or "").strip().lower()
        if type_v not in ("buy", "sell"):
            skipped[f"unknown_type={type_v!r}"] = skipped.get(
                f"unknown_type={type_v!r}", 0
            ) + 1
            continue

        volume = _safe_float(col(row, "volume"))
        executed = _safe_float(col(row, "price"))
        if volume is None or executed is None or volume <= 0:
            skipped["bad_numeric"] = skipped.get("bad_numeric", 0) + 1
            continue

        order_id = col(row, "order")
        if order_id is None:
            skipped["missing_order_id"] = skipped.get("missing_order_id", 0) + 1
            continue

        in_deals.append({
            "time": time_val,
            "deal_id": col(row, "deal"),
            "order_id": order_id,
            "symbol": col(row, "symbol"),
            "type": type_v,
            "volume": volume,
            "executed_price": executed,
            "comment": str(col(row, "comment") or "") if "comment" in col_map else "",
        })

    return {
        "in_deals": in_deals,
        "skipped": skipped,
        "total_rows": total_rows,
        "header_row": header_idx,
        "sheet_name": chosen_ws.title,
    }
